import sys,io,zipfile
sys.path.insert(0,sys.argv[1] if len(sys.argv)>1 else '.')
from fw import *
import docx
def run(profile):
  r=Run('documents',profile); print('#### documents',profile)
  with sync_playwright() as pw:
    r.start(pw); pg=r.pg
    t0=time.time(); ok=r.open_via_chooser('#startOpen','audit.docx')
    pg.wait_for_function("()=>document.querySelector('.page-content')&&document.querySelector('.page-content').innerText.includes('Audit Heading One')",timeout=15000)
    r.rec('open DOCX via start Open + picker','PASS' if ok else 'FAIL',f'chooser={ok} {time.time()-t0:.2f}s title={r.title()}')
    r.rec('dirty after open','PASS' if r.dirty() is False else 'FAIL',r.dirty())
    # navigation panel
    pg.click('#contextBtn'); pg.wait_for_timeout(500); nav=pg.evaluate("()=>{const d=[...document.querySelectorAll('.context-drawer,.drawer')].find(e=>e.getBoundingClientRect().width>0&&getComputedStyle(e).visibility!=='hidden'&&e.innerText.trim());return d?d.innerText.replace(/\\s+/g,' ').slice(0,200):null}")
    r.rec('document panel (outline/navigation) opens','PASS' if nav and 'Second Section' in nav else 'WARNING',nav)
    pg.keyboard.press('Escape'); pg.wait_for_timeout(300)
    # edit: click at end of first paragraph and type
    para=pg.locator('.page-content').get_by_text('Plain paragraph with').first
    para.click(); pg.keyboard.press('End'); pg.keyboard.type(' EDITED'); pg.wait_for_timeout(400)
    has=pg.evaluate("()=>document.querySelector('.page-content').innerText.includes('EDITED')")
    r.rec('type text into document','PASS' if has else 'FAIL',has)
    r.rec('dirty indicator after edit','PASS' if r.dirty() else 'FAIL',r.dirty())
    # bold via select word + toolbar
    pg.keyboard.down('Shift'); [pg.keyboard.press('ArrowLeft') for _ in range(6)]; pg.keyboard.up('Shift')
    pg.get_by_role('button',name='Bold').first.click(); pg.wait_for_timeout(300)
    bold=pg.evaluate("()=>{const s=getSelection();if(!s.rangeCount)return null;let n=s.anchorNode;n=n.nodeType===3?n.parentElement:n;return getComputedStyle(n).fontWeight}")
    r.rec('Bold toolbar on selection','PASS' if bold and int(bold)>=600 else 'FAIL',f'fontWeight={bold}')
    # undo / redo
    pg.click('#undoBtn'); pg.wait_for_timeout(300); b2=pg.evaluate("()=>{const s=getSelection();let n=s.anchorNode;if(!n)return null;n=n.nodeType===3?n.parentElement:n;return getComputedStyle(n).fontWeight}")
    pg.click('#redoBtn'); pg.wait_for_timeout(300); b3=pg.evaluate("()=>{const s=getSelection();let n=s.anchorNode;if(!n)return null;n=n.nodeType===3?n.parentElement:n;return getComputedStyle(n).fontWeight}")
    txt_after=pg.evaluate("()=>{const w=document.createTreeWalker(document.querySelector('.page-content'),NodeFilter.SHOW_TEXT);let n;while(n=w.nextNode()){if(n.data.includes('EDITED'))return getComputedStyle(n.parentElement).fontWeight}return null}")
    r.rec('Undo then Redo of bold','PASS' if (b2 and int(b2)<600) and txt_after and int(txt_after)>=600 else 'WARNING',f'afterUndo(sel)={b2} afterRedo(sel)={b3} EDITED-weight-after-redo={txt_after}')
    # zoom
    pg.click('#zoomMenuBtn'); pg.wait_for_timeout(300); z=r.pg.evaluate("()=>[...document.querySelectorAll('[role=menuitem],[role=menuitemradio],.zoom-option,button')].filter(b=>/^\\d+%$|fit|width|page/i.test(b.innerText.trim())&&b.getBoundingClientRect().width>0).map(b=>b.innerText.trim())")
    r.rec('zoom menu opens with options','PASS' if z else 'FAIL',z[:12])
    if z:
      target=[x for x in z if x.startswith('75')] or z[:1]
      pg.get_by_text(target[0],exact=True).last.click(); pg.wait_for_timeout(400)
      zl=pg.locator('#zoomMenuBtn').inner_text(); r.rec('zoom change applied','PASS' if target[0].split('%')[0] in zl else 'WARNING',f'chose {target[0]} label={zl}')
    # insert table -> native prompt cancel, then accept
    before=pg.evaluate("()=>document.querySelectorAll('.page-content table').length"); nd=len(r.dialogs)
    r.dialog_policy='dismiss'; pg.click('#tableBtn'); pg.wait_for_timeout(500)
    after=pg.evaluate("()=>document.querySelectorAll('.page-content table').length")
    r.rec('Insert table → Cancel leaves document unchanged','PASS' if after==before else 'FAIL',f'native dialogs={r.dialogs[nd:]} tables {before}->{after}')
    r.dialog_policy='accept'; pg.locator('.page-content').get_by_text('Bullet B').click(); pg.keyboard.press('End'); pg.click('#tableBtn'); pg.wait_for_timeout(600); r.dialog_policy='dismiss'
    after2=pg.evaluate("()=>document.querySelectorAll('.page-content table').length")
    r.rec('Insert table → OK inserts table (uses native window.prompt)','PASS' if after2==before+1 else 'FAIL',f'tables {before}->{after2}; dialogs={[d["type"] for d in r.dialogs[nd:]]}')
    # save cancel
    r.menu(); pg.click('#saveMenuBtn'); pg.wait_for_timeout(500); m=r.modal()
    c=r.click_modal(r'^cancel'); n0=len(r.saved())
    r.rec('Save copy panel → Cancel','PASS' if m and c and n0==0 and r.dirty() else 'FAIL',f'panel={m and m["text"][:120]} saved={n0} dirty={r.dirty()}')
    # save with picker abort
    if profile=='fs':
      r.setmode('abort'); r.menu(); pg.click('#saveMenuBtn'); pg.wait_for_timeout(400); r.click_modal(r'save'); pg.wait_for_timeout(700)
      m=r.modal(); r.rec('Save → system picker cancelled (AbortError)','PASS' if r.dirty() and len(r.saved())==0 else 'FAIL',f'dirty={r.dirty()} modal={m and m["text"][:150]}'); r.clear()
      r.setmode('writefail'); r.menu(); pg.click('#saveMenuBtn'); pg.wait_for_timeout(400); r.click_modal(r'save'); pg.wait_for_timeout(900)
      m=r.modal(); st=pg.evaluate("()=>document.querySelector('.statusbar')?.innerText")
      r.rec('Save → write error surfaced and doc stays dirty','PASS' if r.dirty() and ((m and re.search('fail|could not|error|not',m['text'],re.I)) or re.search('fail|could not|error',st or '',re.I)) else 'FAIL',f'dirty={r.dirty()} modal={m and m["text"][:160]} status={st}'); shot=r.shot('writefail'); r.clear()
      r.setmode('ok')
    t1=time.time(); r.menu(); pg.click('#saveMenuBtn'); pg.wait_for_timeout(400); r.click_modal(r'save'); pg.wait_for_timeout(1500)
    if profile=='fs':
      sv=r.saved(); dt=time.time()-t1
      ok=False; det=''
      if sv:
        data=r.saved_bytes(); d=docx.Document(io.BytesIO(data)); txt='\n'.join(p.text for p in d.paragraphs); ok='EDITED' in txt and 'Audit Heading One' in txt and len(d.tables)==1
        bolds=[run.text for p in d.paragraphs for run in p.runs if run.bold]
        det=f'{sv[-1]} tables={len(d.tables)} bold_runs={bolds} {dt:.2f}s'
        open(f'{S}/func/doc_saved.docx','wb').write(data)
      r.rec('Save copy → DOCX written, valid, contains edits','PASS' if ok else 'FAIL',det)
      r.rec('dirty cleared after confirmed save','PASS' if r.dirty() is False else 'WARNING',r.dirty())
    else:
      pg.wait_for_timeout(800); r.rec('Save copy (no FS API) → download fallback','PASS' if r.downloads else 'FAIL',[d['name'] for d in r.downloads]); m=r.modal(); r.rec('post-download state',"INFO",f'dirty={r.dirty()} modal={m and m["text"][:200]}'); r.click_modal(r'done|close|ok|keep|cancel')
    # reopen saved
    if profile=='fs' and os.path.exists(f'{S}/func/doc_saved.docx'):
      import shutil; shutil.copy(f'{S}/func/doc_saved.docx',f'{FX}/doc_saved.docx')
      r.menu(); ok=r.open_via_chooser('#openMenuBtn','doc_saved.docx'); pg.wait_for_timeout(1500); m=r.modal()
      if m: r.rec('guard on open while clean?','INFO',m['text'][:150]); r.click_modal(r'discard|open|continue|don')
      has=pg.evaluate("()=>document.querySelector('.page-content').innerText.includes('EDITED')")
      r.rec('Reopen saved DOCX shows edits','PASS' if has else 'FAIL',f'chooser={ok} title={r.title()}')
    # dirty guard when opening another file
    para=pg.locator('.page-content p, .page-content div').first; pg.locator('.page-content').click(); pg.keyboard.type('Z'); pg.wait_for_timeout(300)
    r.menu(); pg.click('#openMenuBtn'); pg.wait_for_timeout(500); m=r.modal()
    r.rec('Open while dirty → unsaved-changes guard shown','PASS' if m else 'FAIL',m and m['text'][:200])
    if m:
      c=r.click_modal(r'^cancel|keep'); r.rec('Guard → Cancel keeps document','PASS' if c and r.dirty() else 'FAIL',f'clicked={c} dirty={r.dirty()}')
      r.menu(); 
      try:
        with pg.expect_file_chooser(timeout=4000) as fc:
          pg.click('#openMenuBtn'); pg.wait_for_timeout(400); r.click_modal(r'discard|don.t save|continue|open')
        fc.value.set_files(f'{FX}/audit.rtf'); pg.wait_for_timeout(1500)
        has=pg.evaluate("()=>document.querySelector('.page-content').innerText.includes('RTF audit line one')")
        r.rec('Guard → Discard then open RTF','PASS' if has else 'FAIL',f'title={r.title()}')
      except Exception as e: r.rec('Guard → Discard then open RTF','FAIL',repr(e)[:200])
    m=r.modal(); 
    if m: r.rec('modal after RTF open','INFO',m['text'][:200]); r.click_modal(r'close|ok|continue|cancel')
    # DOC legacy
    r.open_via_input('audit.doc'); pg.wait_for_timeout(2500); m=r.modal()
    if m: r.rec('DOC open shows notice/guard','INFO',m['text'][:200]); r.click_modal(r'discard|continue|open|ok|close')
    pg.wait_for_timeout(1500)
    has=pg.evaluate("()=>document.querySelector('.page-content').innerText.includes('Audit Heading One')"); r.rec('Open legacy DOC → content visible','PASS' if has else 'FAIL',f'title={r.title()}')
    r.shot('doc')
    # invalid & empty
    for bad in ['invalid.docx','empty.docx','invalid.doc']:
      r.open_via_input(bad); pg.wait_for_timeout(1500); m=r.modal()
      if m and re.search('discard|unsaved',m['text'],re.I): r.click_modal(r'discard|continue|don'); pg.wait_for_timeout(1200); m=r.modal()
      st=pg.evaluate("()=>document.querySelector('.statusbar')?.innerText"); ttl=r.title()
      r.rec(f'Open {bad} → clear error, previous doc preserved','PASS' if (m and re.search('could not|unsupported|invalid|error|not a|damaged|corrupt|empty',m['text'],re.I)) else 'FAIL',f'modal={m and m["text"][:200]} status={st} title={ttl}')
      r.click_modal(r'close|ok|dismiss|cancel|back')
    usable=pg.evaluate("()=>document.querySelector('.page-content')?.innerText.length>0"); r.rec('App still usable after invalid files','PASS' if usable else 'FAIL',r.title())
    # new document
    r.menu(); pg.click('#newMenuBtn'); pg.wait_for_timeout(600); m=r.modal()
    if m: r.click_modal(r'discard|continue|don|new'); pg.wait_for_timeout(600)
    r.rec('New document','PASS' if 'Untitled' in (pg.locator('#titleText').input_value() if pg.locator('#titleText').evaluate("e=>e.tagName")=='INPUT' else pg.locator('#titleText').inner_text()) else 'WARNING',r.title())
    # unload guard
    pg.locator('.page-content').click(); pg.keyboard.type('unsaved'); pg.wait_for_timeout(300)
    r.rec('beforeunload guard when dirty','PASS' if r.before_unload_guard() else 'FAIL','')
    r.finish()
  return r
if __name__=='__main__':
  for prof in (sys.argv[2:] or ['fs','nofs']): run(prof)
