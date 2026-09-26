import docx, openpyxl, pptx, zipfile, io, os
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
# DOCX
d=docx.Document(); d.add_heading('Audit Heading One',1); p=d.add_paragraph('Plain paragraph with '); r=p.add_run('bold'); r.bold=True; p.add_run(' and '); r=p.add_run('italic'); r.italic=True; p.add_run(' text.')
d.add_heading('Second Section',2); d.add_paragraph('Bullet A',style='List Bullet'); d.add_paragraph('Bullet B',style='List Bullet')
t=d.add_table(rows=2,cols=2); t.cell(0,0).text='A1'; t.cell(1,1).text='B2'; d.save('audit.docx')
open('audit.rtf','w').write(r'{\rtf1\ansi\deff0{\fonttbl{\f0 Times New Roman;}}\f0\fs24 RTF audit line one \b bold\b0  and \i italic\i0 .\par Second RTF paragraph.\par}')
# XLSX
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Data'
for i in range(1,11): ws.cell(i,1,i); ws.cell(i,2,f'Item {i}'); ws.cell(i,3,f'=A{i}*2')
ws['E1']='=SUM(A1:A10)'; ws['E2']='00123'; ws['E2'].number_format='@'; wb.create_sheet('Second')['A1']='sheet2'; wb.save('audit.xlsx')
# PPTX
pr=pptx.Presentation()
for i in range(3):
  s=pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text=f'Slide {i+1} Title'; s.placeholders[1].text=f'Body text for slide {i+1}'
  s.shapes.add_shape(1,Inches(6),Inches(4),Inches(2),Inches(1)).text='Box'
pr.save('audit.pptx')
open('audit.txt','w').write('Line one of plain text\nLine two\n\nLast line with ümlaut and 中文\n')
# EPUB
def epub(name,title,chapters):
  z=zipfile.ZipFile(name,'w'); z.writestr(zipfile.ZipInfo('mimetype'),'application/epub+zip',compress_type=zipfile.ZIP_STORED)
  z.writestr('META-INF/container.xml','<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',compress_type=zipfile.ZIP_DEFLATED)
  man=''.join(f'<item id="c{i}" href="c{i}.xhtml" media-type="application/xhtml+xml"/>' for i in range(len(chapters)))
  spine=''.join(f'<itemref idref="c{i}"/>' for i in range(len(chapters)))
  z.writestr('OEBPS/content.opf',f'<?xml version="1.0" encoding="UTF-8"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">urn:uuid:{title}</dc:identifier><dc:title>{title}</dc:title><dc:language>en</dc:language><meta property="dcterms:modified">2026-01-01T00:00:00Z</meta></metadata><manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>{man}</manifest><spine>{spine}</spine></package>')
  nav=''.join(f'<li><a href="c{i}.xhtml">{c[0]}</a></li>' for i,c in enumerate(chapters))
  z.writestr('OEBPS/nav.xhtml',f'<?xml version="1.0" encoding="UTF-8"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>nav</title></head><body><nav epub:type="toc"><ol>{nav}</ol></nav></body></html>')
  for i,(h,body) in enumerate(chapters):
    z.writestr(f'OEBPS/c{i}.xhtml',f'<?xml version="1.0" encoding="UTF-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>{h}</title></head><body><h1>{h}</h1>'+''.join(f'<p>{body} paragraph {j}.</p>' for j in range(40))+'</body></html>')
  z.close()
epub('audit.epub','Audit Book',[('Chapter One','Alpha'),('Chapter Two','Beta'),('Chapter Three','Gamma')])
epub('audit2.epub','Second Book',[('Other Start','Delta'),('Other End','Epsilon')])
# PDFs
def pdf(name,pages,title):
  c=canvas.Canvas(name,pagesize=A4)
  for i in range(pages):
    c.setFont('Helvetica-Bold',28); c.drawString(72,760,f'{title} page {i+1}')
    c.setFont('Helvetica',11)
    for j in range(40): c.drawString(72,720-j*15,f'Line {j} of page {i+1} lorem ipsum dolor sit amet consectetur')
    c.showPage()
  c.save()
pdf('small.pdf',3,'Small PDF'); pdf('second.pdf',2,'Second PDF')
