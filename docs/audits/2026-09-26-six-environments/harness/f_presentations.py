import sys,io,shutil
sys.path.insert(0,sys.argv[1]); from fw import *
import pptx
def I(pg): return pg.evaluate("()=>__inkdosPresentations.inspect().session")
def canvas_text(pg): return pg.evaluate("()=>{const c=[...document.querySelectorAll('[class*=slide-canvas],[class*=slide-stage],.slide-surface')].find(e=>e.getBoundingClientRect().left>200);return c?c.innerText:document.body.innerText}")
def run(profile):
  r=Run('presentations',profile); print('#### presentations',profile)
  with sync_playwright() as pw:
    r.start(pw); pg=r.pg
    ok=r.open_via_chooser('#startOpen','audit.pptx'); pg.wait_for_timeout(2000); s=I(pg)
    r.rec('open PPTX via start Open + picker','PASS' if ok and s['slideCount']==3 else 'FAIL',s)
    th=pg.evaluate("()=>document.querySelectorAll('#slidePanel .slide-thumb,[class*=thumb]').length"); r.rec('thumbnails rendered','PASS' if th>=3 else 'FAIL',th)
    pg.click('#nextSlideBtn'); pg.wait_for_timeout(300); a=I(pg)['currentIndex']; pg.keyboard.press('Escape'); pg.click('#prevSlideBtn'); pg.wait_for_timeout(300); b=I(pg)['currentIndex']
    r.rec('next/previous slide','PASS' if a==1 and b==0 else 'FAIL',f'{a},{b}')
    # select + edit text
    t=pg.locator('span',has_text='Body text for slide 1').last; t.click(); pg.wait_for_timeout(300); sel=I(pg)['selectedObjectId']
    r.rec('click selects object','PASS' if sel else 'FAIL',sel)
    t.dblclick(); pg.wait_for_timeout(300); pg.keyboard.press('End'); pg.keyboard.type(' EDITED'); pg.wait_for_timeout(200); pg.keyboard.press('Escape'); pg.wait_for_timeout(400)
    r.rec('edit text in placeholder','PASS' if 'Body text for slide 1 EDITED' in canvas_text(pg) else 'FAIL',canvas_text(pg)[:120].replace('\n',' | '))
    r.rec('dirty after edit','PASS' if r.dirty() and I(pg)['dirty'] else 'FAIL',I(pg)['dirty'])
    # bold on selected object
    t=pg.locator('span',has_text='Slide 1 Title').last; t.click(); pg.wait_for_timeout(200)
    fb=pg.evaluate("()=>[...document.querySelectorAll('button')].filter(b=>/^bold$/i.test(b.getAttribute('aria-label')||b.title||'')&&b.getBoundingClientRect().width>0).map(b=>b.id||b.className)")
    r.rec('contextual text toolbar appears on selection (Bold)','PASS' if fb else 'WARNING',fb)
    # add/duplicate/delete/move + undo
    n0=I(pg)['slideCount']; pg.click('#addSlideBtn'); pg.wait_for_timeout(400); n1=I(pg)['slideCount']
    pg.click('#duplicateSlideBtn'); pg.wait_for_timeout(400); n2=I(pg)['slideCount']
    pg.click('#deleteSlideBtn'); pg.wait_for_timeout(400); m=r.modal()
    if m: r.rec('delete slide confirmation','INFO',m['text'][:100]); r.click_modal(r'delete|ok|confirm')
    n3=I(pg)['slideCount']; pg.click('#undoBtn'); pg.wait_for_timeout(400); n4=I(pg)['slideCount']
    r.rec('add / duplicate / delete slide / undo delete','PASS' if (n0,n1,n2,n3,n4)==(3,4,5,4,5) else 'FAIL',(n0,n1,n2,n3,n4))
    pg.click('#redoBtn'); pg.wait_for_timeout(300); n5=I(pg)['slideCount']; r.rec('redo','PASS' if n5==4 else 'FAIL',n5)
    idx=I(pg)['currentIndex']; pg.click('#moveSlideUpBtn'); pg.wait_for_timeout(300); idx2=I(pg)['currentIndex']; r.rec('move slide earlier','PASS' if idx2==idx-1 else 'FAIL',f'{idx}->{idx2}')
    # insert text box, shape
    before=pg.evaluate("()=>__inkdosPresentations.session.current?.objects?.length ?? null")
    pg.click('#insertTextBtn'); pg.wait_for_timeout(400); r.rec('insert text box','PASS' if I(pg)['selectedObjectId'] else 'WARNING',I(pg)['selectedObjectId']); pg.keyboard.press('Escape')
    # background
    pg.click('#pptP1Background'); pg.wait_for_timeout(400); m=r.modal(); pop=pg.evaluate("()=>[...document.querySelectorAll('[class*=popover],[class*=menu],[role=dialog]')].filter(e=>e.getBoundingClientRect().width>0&&getComputedStyle(e).visibility!=='hidden').map(e=>e.id||e.className).slice(0,3)")
    r.rec('slide background control opens','PASS' if (m or pop) else 'FAIL',f'modal={m and m["text"][:80]} pop={pop}'); r.shot('background'); pg.keyboard.press('Escape'); r.clear()
    # zoom
    pg.click('#zoomMenuBtn'); pg.wait_for_timeout(300); z=pg.evaluate("()=>[...document.querySelectorAll('button,[role=menuitemradio]')].filter(b=>/^\\d+%$|fit/i.test(b.innerText.trim())&&b.getBoundingClientRect().width>0).map(b=>b.innerText.trim())")
    r.rec('zoom menu','PASS' if z else 'FAIL',z); pg.keyboard.press('Escape')
    # present
    pg.click('#presentBtn'); pg.wait_for_timeout(900); pres=pg.evaluate("()=>!!document.fullscreenElement||[...document.querySelectorAll('[class*=slideshow],[class*=present]')].some(e=>e.getBoundingClientRect().width>1000)")
    r.shot('present'); pg.keyboard.press('ArrowRight'); pg.wait_for_timeout(300); pg.keyboard.press('Escape'); pg.wait_for_timeout(600)
    back=pg.evaluate("()=>![...document.querySelectorAll('[class*=slideshow]')].some(e=>e.getBoundingClientRect().width>1000&&getComputedStyle(e).display!=='none')")
    r.rec('Present mode opens and Escape exits','PASS' if pres and back else 'FAIL',f'opened={pres} exited={back}')
    # save
    def do_save():
      r.menu(); pg.click('#saveMenuBtn'); pg.wait_for_timeout(700); m=r.modal()
      if m: r.click_modal(r'save|copy|pptx|continue'); pg.wait_for_timeout(1500)
      return m
    if profile=='fs':
      r.setmode('abort'); do_save(); r.rec('Save → picker cancelled keeps dirty','PASS' if r.dirty() and not r.saved() else 'FAIL',r.dirty()); r.clear()
      r.setmode('writefail'); do_save(); m=r.modal(); r.rec('Save → write error surfaced','PASS' if r.dirty() and m and re.search('could not|fail|error',m['text'],re.I) else 'FAIL',m and m['text'][:200]); r.shot('writefail'); r.clear()
      r.setmode('ok'); do_save(); sv=r.saved()
      if sv:
        data=r.saved_bytes(); open(f'{S}/func/ppt_saved.pptx','wb').write(data)
        try:
          P=pptx.Presentation(io.BytesIO(data)); texts=[' '.join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in P.slides]
          r.rec('Save → PPTX valid, slide count & edit preserved','PASS' if len(P.slides)==I(pg)['slideCount'] and any('EDITED' in t for t in texts) else 'FAIL',f'slides={len(P.slides)} texts={[t[:40] for t in texts]}')
        except Exception as e: r.rec('Save → PPTX loads in python-pptx','FAIL',repr(e)[:300])
      else: r.rec('Save → PPTX written','FAIL',sv)
      r.rec('dirty cleared after save','PASS' if not r.dirty() else 'WARNING',r.dirty()); r.clear()
      shutil.copy(f'{S}/func/ppt_saved.pptx',f'{FX}/ppt_saved.pptx')
      r.menu(); r.open_via_chooser('#openMenuBtn','ppt_saved.pptx'); pg.wait_for_timeout(2000); m=r.modal()
      if m: r.click_modal(r'discard|open|continue'); pg.wait_for_timeout(1500)
      found=pg.evaluate("()=>document.body.innerText.includes('EDITED')"); r.rec('reopen saved PPTX shows edit','PASS' if found else 'FAIL',f'{I(pg)} ')
    else:
      do_save(); pg.wait_for_timeout(800); r.rec('Save (no FS API) → download fallback','PASS' if r.downloads else 'FAIL',[d['name'] for d in r.downloads]); r.clear()
    # guard
    pg.click('#addSlideBtn'); pg.wait_for_timeout(300)
    r.menu(); pg.click('#openMenuBtn'); pg.wait_for_timeout(500); m=r.modal(); r.rec('Open while dirty → guard','PASS' if m else 'FAIL',m and m['text'][:150])
    covered=pg.evaluate("()=>{const d=document.getElementById('presentationsUnsavedDialog');if(!d)return null;const cs=getComputedStyle(d);const b=d.querySelector('[data-choice=discard]').getBoundingClientRect();const hit=document.elementFromPoint(b.x+b.width/2,b.y+b.height/2);return {position:cs.position,zIndex:cs.zIndex,bg:cs.backgroundColor,discardHit:hit&&(hit.id||hit.className),discardY:Math.round(b.y)}}")
    r.rec('Presentations unsaved guard is styled/overlaid and reachable from menu','FAIL' if covered and (covered['position']=='static' or covered['discardHit']!='') else 'PASS',covered)
    if m:
      pg.keyboard.press('Escape'); pg.wait_for_timeout(300)
      pg.locator('#presentationsUnsavedDialog [data-choice=cancel]').click(); pg.wait_for_timeout(300)
      r.rec('Guard → Cancel keeps presentation','PASS' if r.dirty() else 'FAIL',r.dirty())
      r.menu(); pg.click('#openMenuBtn'); pg.wait_for_timeout(400); pg.keyboard.press('Escape'); pg.wait_for_timeout(300)
      try:
        with pg.expect_file_chooser(timeout=4000) as fc:
          pg.locator('#presentationsUnsavedDialog [data-choice=discard]').click()
        fc.value.set_files(f'{FX}/audit.ppt'); pg.wait_for_timeout(3500); m=r.modal()
        if m: r.rec('after PPT chosen modal','INFO',m['text'][:160]); r.clear()
        s=I(pg); r.rec('Guard → Discard → open legacy PPT','PASS' if s['slideCount']==3 and 'Slide 1 Title' in pg.evaluate("()=>document.body.innerText") else 'FAIL',s)
        r.shot('ppt')
      except Exception as e: r.rec('Guard → Discard → open PPT (menu closed first)','FAIL',repr(e)[:200])
    for bad in ['invalid.pptx']:
      r.open_via_input(bad); pg.wait_for_timeout(1500); m=r.modal()
      if m and re.search('unsaved|discard',m['text'],re.I): r.click_modal(r'^discard'); pg.wait_for_timeout(1200); m=r.modal()
      r.rec(f'Open {bad} → error, previous kept','PASS' if m and re.search('could not|invalid|error|not',m['text'],re.I) and I(pg)['slideCount']>0 else 'FAIL',f'modal={m and m["text"][:200]} s={I(pg)}'); r.clear()
    shutil.copy(f'{FX}/empty.txt',f'{FX}/empty.pptx'); r.open_via_input('empty.pptx'); pg.wait_for_timeout(1500); m=r.modal()
    if m and re.search('unsaved|discard',m['text'],re.I): r.click_modal(r'^discard'); pg.wait_for_timeout(1200); m=r.modal()
    r.rec('Open empty.pptx → error','PASS' if m and re.search('could not|invalid|error|not|empty',m['text'],re.I) else 'FAIL',m and m['text'][:200]); r.clear()
    r.rec('legacy PPT opens read-only (toolbar disabled, "Legacy PPT is read-only")','INFO',pg.evaluate("()=>document.getElementById('addSlideBtn').title"))
    r.open_via_input('audit.pptx'); pg.wait_for_timeout(2000); r.clear(); pg.click('#addSlideBtn')
    pg.wait_for_timeout(300)
    r.rec('beforeunload guard when dirty','PASS' if r.before_unload_guard() else 'FAIL','')
    r.finish()
for prof in sys.argv[2:]: run(prof)
